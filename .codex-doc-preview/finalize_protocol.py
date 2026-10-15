from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

f=Path(r'C:\Users\euco\Desktop\PROYECTO\front\deliverables\PROTOCOLO_Proyecto_Grado_2026_actualizado_con_evidencia.docx')
d=Document(f)

# Final scope consistency updates based on the user's confirmation.
for p in d.paragraphs:
    s=p.text.strip()
    if s.startswith('El borrador consigna 5.000–6.000 ingresos diarios'):
        p.text='La población de vehículos y el flujo de ingreso al campus no se estiman en esta versión: no se realiza trabajo de campo y no se aportó una fuente institucional para las cifras del borrador. Se eliminaron los 359 registros del diseño muestral porque no representan imágenes del benchmark ni datos recogidos. Un trabajo futuro con campo requeriría definir población, periodo, unidad y muestra a partir de registros verificables.'
    elif s.startswith('Definir un protocolo para una eventual prueba piloto'):
        p.text='Evaluar técnicamente el detector YOLOv8n en el split test reservado y documentar métricas, configuración, limitaciones y evidencia reproducible. La evaluación con cámaras del campus queda fuera del alcance de esta versión.'
    elif s.startswith('A partir de los resultados y límites documentados'):
        p.text='A partir de los resultados y límites documentados, se recomiendan las siguientes mejoras técnicas y líneas de trabajo futuro:'
    elif s.startswith('Completar la comparación de tiempos y errores') or s.startswith('Verificar N, periodo y unidad de análisis'):
        p.text='Una evaluación futura en campus requeriría autorización institucional, fuente verificable de población, periodo y unidad de análisis, protocolo de privacidad y diseño muestral reproducible. Esta actividad está fuera del alcance actual.'
    elif s.startswith('Caracterizar el control de acceso mediante revisión documental'):
        p.text='Describir la problemática y necesidades del acceso vehicular a partir de revisión documental disponible, distinguiendo antecedentes reportados de datos institucionales verificados. No se aplicarán encuestas ni observación de campo en esta versión.'
    elif s.startswith('Describir la problemática y necesidades del acceso vehicular'):
        p.text='Documentar el problema y los requisitos funcionales del prototipo mediante revisión del protocolo y de la documentación técnica disponible, diferenciando antecedentes reportados de datos verificados. No se aplicarán encuestas ni observación de campo en esta versión.'
    elif s.startswith('Evaluar técnicamente el detector YOLOv8n en el split test reservado'):
        p.text='Evaluar técnicamente el detector YOLOv8n sobre el split test reservado y documentar métricas, configuración, límites y evidencia reproducible. La evaluación con cámaras del campus queda fuera del alcance de esta versión.'
    elif s.startswith('Actualmente, el sistema manual de registro de motocicletas'):
        p.text='El protocolo describe un procedimiento manual que depende de la supervisión del personal y de registros físicos. Las posibles demoras, congestión, errores y pérdidas de información son aspectos por caracterizar; no se aportó una línea base para cuantificar su frecuencia o impacto.'
    elif s.startswith('La necesidad de fortalecer y actualizar los procesos de control de acceso'):
        p.text='La motivación del proyecto es explorar una herramienta tecnológica para apoyar el control de acceso vehicular descrito en el protocolo. La magnitud de sus efectos sobre seguridad, movilidad y eficiencia no se midió en esta versión.'
    elif s.startswith('Desde el punto de vista práctico, el proyecto permitirá eliminar'):
        p.text='Desde el punto de vista práctico, el prototipo explora el registro digital de movimientos. La sustitución de papel, la trazabilidad completa, las consultas inmediatas y cualquier ahorro operativo dependen de integración y despliegue institucional que no están verificados en esta versión.'
    elif s.startswith('Más allá del beneficio que aportará a la institución'):
        p.text='El prototipo permite explorar una alternativa para el control de acceso. No se midió reducción de riesgos, inconsistencias ni brechas de seguridad, y su escalabilidad a otras instituciones requiere evaluación independiente.'
    elif s.startswith('Académicamente, esta investigación aporta al fortalecimiento'):
        p.text='En el ámbito académico, el trabajo documenta una aplicación de detección de placas con visión artificial. La replicabilidad en otras instituciones dependerá de sus datos, condiciones de captura, infraestructura y reglas de operación.'
    elif s.startswith('el diagnóstico del problema actual de acceso vehicular permite medir'):
        p.text='El trabajo documenta el prototipo y evalúa la localización de placas en un conjunto test reservado. No se realizaron diagnóstico cuantitativo de campo, mediciones de ingreso, encuestas de satisfacción ni implementación en condiciones reales; esos resultados no se afirman.'
    elif s.startswith('En consecuencia, este proyecto no solo busca atender'):
        p.text='En consecuencia, esta versión aporta un prototipo técnico y resultados de evaluación acotados al conjunto de imágenes utilizado. Su uso operativo futuro requerirá integración, validación y aprobación institucional.'
    elif s.startswith('La implementación de un sistema de acceso vehicular mediante reconocimiento automático'):
        p.text='La construcción del prototipo explora una alternativa para el proceso manual descrito en el documento. Esta versión no prueba que la solución reduzca congestión, tiempos o errores en la Universidad Popular del Cesar.'
    elif s.startswith('En primer lugar, se busca lograr una mejora significativa'):
        p.text='El prototipo tiene como motivación apoyar el registro durante periodos de alta demanda. No se midieron los tiempos actuales ni los del sistema; por ello, no se afirma que la automatización reduzca esperas o mejore el flujo vehicular.'
    elif s.startswith('En segundo lugar, se contempla una reducción considerable'):
        p.text='La lectura asistida por software podría reducir digitación manual, pero no se cuantificó la tasa de errores ni se midió una reducción. La exactitud OCR tampoco se estableció con una muestra representativa.'
    elif s.startswith('Asimismo, el sistema permitirá una mayor trazabilidad institucional'):
        p.text='El prototipo contempla registrar movimientos digitalmente. La persistencia completa, consulta histórica y efecto en seguridad requieren verificar integración con SQL Server y operación autorizada; no se presentan como resultados demostrados.'
    elif s.startswith('Finalmente, se estima una alta escalabilidad'):
        p.text='La escalabilidad y los costos de operación no se evaluaron. Su estimación exigiría definir infraestructura, volumen, disponibilidad, mantenimiento e integración con cámaras institucionales.'
    elif s.startswith('Por tanto, la aplicación práctica de este proyecto'):
        p.text='La aplicación práctica de esta versión consiste en un prototipo técnico y una evaluación del detector sobre un benchmark. La adopción institucional y la automatización de procesos quedan para una decisión y evaluación posteriores.'
    elif s.startswith('The project develops a web prototype to support motorcycle access control'):
        p.text='Abstract – The project develops a web prototype to support motorcycle access control at Universidad Popular del Cesar using YOLO plate detection. This version has a technical scope and does not include a campus field pilot. The draft mentions 5,000–6,000 motorcycles per day, but no institutional source was provided; the figure remains unverified and is not used to calculate a study sample. The fine-tuned YOLOv8n detector was evaluated on the reserved test split of an external dataset: 101 images, including 67 annotated plates and 34 background images. It achieved 95.6% detection precision, 96.8% recall, 99.0% mAP@50, and 88.8% mAP@50–95. These metrics measure box localization, not OCR transcription. Exact OCR evaluation, end-to-end frontend–backend–SQL integration, and institutional deployment are not established by the available evidence.'
    elif s.startswith('Keywords –'):
        p.text=s.replace('computer visión','computer vision')
    elif s.startswith('En el contexto universitario, la integración de estas tecnologías permite fortalecer'):
        p.text='En el contexto universitario, estas tecnologías podrían apoyar el control de acceso; los beneficios concretos dependen de su evaluación operativa e integración en cada institución.'
    elif s.startswith('La investigación se desarrollará bajo un diseño de desarrollo tecnológico'):
        p.text='El diseño corresponde al desarrollo de un prototipo tecnológico y a su evaluación cuantitativa sobre un conjunto de imágenes etiquetadas. No contempla una fase de operación en porterías.'
    elif s.startswith('Fase 1. Diagnóstico y levantamiento:'):
        p.text='Fase 1. Diagnóstico documental: el protocolo caracteriza el problema a partir de antecedentes reportados. No se adjuntaron actas de observación, encuestas, entrevistas, cronometrajes ni registros institucionales de afluencia; por tanto, no se afirma que se haya levantado una línea base de campo.'
    elif s.startswith('La discusión se limita a las métricas de detección'):
        p.text='La discusión se limita a las métricas de detección de placas en el split test y a describir el pequeño smoke test OCR como exploratorio. No se afirman exactitud OCR general, reducción de tiempos o errores, congestión, aceptación de usuarios ni rendimiento en cámaras del campus. Una evaluación de esos aspectos sería un estudio futuro con datos y permisos adecuados.'
    elif s.startswith('La solución contempla un backend Microsoft ASP.NET Core'):
        p.text='La solución contempla un backend Microsoft ASP.NET Core con patrón Endpoint/REPR de Ardalis, persistencia prevista en SQL Server mediante Dapper y documentación con Swagger; Azure es la infraestructura propuesta. La evidencia reunida corresponde a un prototipo y una prueba temporal local del servicio de visión; no acredita despliegue en Azure ni una prueba integral sostenida entre interfaz, backend y base de datos. La reducción de tiempos y errores y cualquier mejora de seguridad serían asuntos de una evaluación futura independiente, fuera del alcance de esta versión.'
    elif s.startswith('La evaluación ejecutada del detector se realizó sobre imágenes etiquetadas'):
        p.text='La evaluación del detector se realizó sobre imágenes etiquetadas, no mediante recolección en el campus. Una evaluación de campo sería un estudio posterior, independiente y fuera del alcance de esta versión.'
    elif s.startswith('El proyecto se encuentra alineado con el plan de transformación digital'):
        p.text='El proyecto se relaciona con los objetivos de transformación digital y fortalecimiento de capacidades tecnológicas descritos en el Plan de Transformación Digital del Ministerio de Educación Nacional [3]. Esta referencia contextual no se interpreta como una obligación normativa para que cada universidad integre inteligencia artificial en todos sus procesos.'
    elif s.startswith('Este proyecto se encuentra alineado con el plan de transformación digital'):
        p.text='El proyecto se relaciona con los objetivos de transformación digital y fortalecimiento de capacidades tecnológicas descritos en el Plan de Transformación Digital del Ministerio de Educación Nacional [3]. Esta referencia contextual no se interpreta como una obligación normativa para que cada universidad integre inteligencia artificial en todos sus procesos.'
    elif s.startswith('De igual forma, estudios relacionados con el preprocesamiento de imágenes'):
        p.text='Los estudios de superresolución, como el de Ledig et al. (2017) [5], analizan la reconstrucción de imágenes mediante redes generativas. Ese trabajo no evalúa placas vehiculares y no demuestra mejoras de OCR en este proyecto; se conserva solo como antecedente general de procesamiento de imágenes.'
    elif s.startswith('Asimismo, estudios desarrollados por Julián et al. (2023)'):
        p.text='Laroca et al. [10] describen un sistema ALPR que integra detección basada en YOLO y etapas de reconocimiento de caracteres. Sus resultados corresponden a sus propios conjuntos de datos y no permiten anticipar el rendimiento del prototipo de la UPC.'
    elif s.startswith('Por otra parte, investigaciones recientes como la de Orozco Calderón'):
        p.text='Orozco Calderón y Zúñiga Rendon [22] presentan un trabajo de titulación sobre identificación de placas mediante visión artificial en tiempo real. Sus resultados y condiciones pertenecen a ese proyecto y no constituyen validación del prototipo aquí desarrollado.'
    elif s.startswith('Su integración con modelos de inteligencia artificial amplía su funcionalidad'):
        p.text='La integración de videovigilancia con visión artificial puede explorarse para detectar eventos; las funciones de alertas, vehículos no autorizados y control de permanencia no están implementadas ni evaluadas en esta versión. Cualquier captura futura deberá someterse a revisión institucional de privacidad [11].'

# Update the schedule on the cover: dates remain the original plan, not proof of completion.
for table in d.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                s=p.text.strip()
                if s.startswith('El presente proyecto será desarrollado en la ciudad de Valledupar'):
                    p.text='El prototipo se desarrolla como trabajo de grado en la Universidad Popular del Cesar, sede Campus Universitario. La evaluación documentada en esta versión es técnica y se realiza sobre imágenes etiquetadas; no incluye implementación ni piloto en las porterías.'
                elif s.startswith('El proyecto tendrá un lapso de ejecución de cinco (5) meses'):
                    p.text='El cronograma inicial programó cinco (5) meses, del 01 de marzo al 31 de julio de 2026. Estas fechas corresponden al plan original y no acreditan por sí solas que cada actividad se haya ejecutado; el estado se describe en la sección 3.2.'
                elif s.startswith('Se realizará la caracterización del proceso actual'):
                    p.text='Fase documental: describir el problema con antecedentes disponibles. No se presentan como recolectados tiempos, encuestas ni cifras de afluencia sin evidencia institucional.'
                elif s.startswith('Se ejecutarán pruebas funcionales, pruebas de precisión'):
                    p.text='Se evaluó técnicamente el detector sobre el split test reservado. No se ejecutaron pruebas operativas en condiciones reales del campus.'
                elif s.startswith('Se realizarán las correcciones necesarias'):
                    p.text='Se documentan las métricas técnicas disponibles y los límites del prototipo. El calendario y cierre administrativo deben actualizarse con las fechas reales de entregables verificables.'
                elif s.startswith('La investigación tendrá aplicación directa en la sede Campus Universitario'):
                    p.text='El sistema se propone para apoyar el control de acceso vehicular en el Campus Universitario; esta versión no fue desplegada ni evaluada en operación institucional.'
                elif s=='Las actividades se desarrollarán en las siguientes fases:':
                    p.text='El plan original organizó las actividades en las siguientes fases:'
                elif s.startswith('El presente proyecto se desarrollará en la ciudad de Valledupar'):
                    p.text='El prototipo se desarrolla como trabajo de grado y se orienta al contexto de la Universidad Popular del Cesar. La evaluación de esta versión usa imágenes etiquetadas y no incluye despliegue en el campus.'

# Repair verifiable reference metadata, remove duplicate Law 1581 entry, and number the added source.
for p in list(d.paragraphs):
    s=p.text.strip()
    if s.startswith('[5].Leidig'):
        p.text='[5] C. Ledig, L. Theis, F. Huszár, J. Caballero, A. Cunningham, A. Acosta, A. Aitken, A. Tejani, J. Totz, Z. Wang, and W. Shi, “Photo-Realistic Single Image Super-Resolution Using a Generative Adversarial Network,” in Proceedings of the IEEE Conference on Computer Vision and Pattern Recognition (CVPR), 2017, pp. 4681–4690.'
    elif s.startswith('[7]. Universidad Estatal Península de Santa Elena'):
        p.text='[7] A. Villacis Jeikol Marcelo, “Detección y registro de placas vehiculares aplicando visión por computadora para acceder a las instalaciones universitarias UPSE,” trabajo de integración curricular, Universidad Estatal Península de Santa Elena, Santa Elena, Ecuador, 2024.'
    elif s.startswith('[9]. Universidad EAFIT'):
        p.text='[9] D. E. Pinto Restrepo, “Sistema de reconocimiento de placas colombianas por medio de redes convolucionales para acceso a áreas residenciales,” tesis de maestría, Universidad EAFIT, Escuela de Ciencias Aplicadas e Ingeniería, Medellín, Colombia, 2024.'
    elif s.startswith('[21] Congreso de la República de Colombia, Ley 1581'):
        p._element.getparent().remove(p._element)
    elif s.startswith('[22] M. P. Groover'):
        p.text='[21] M. P. Groover, Automation, Production Systems, and Computer-Integrated Manufacturing, 2nd ed. Upper Saddle River, NJ: Prentice Hall, 2001.'
    elif s.startswith('[21] Congreso de Colombia. (2012). Ley 1581'):
        p._element.getparent().remove(p._element)

for p in d.paragraphs:
    if p.text.strip().startswith('Su integración con modelos de inteligencia artificial amplía su funcionalidad'):
        p.text=p.text.replace('[21]','[11]')
    elif p.text.strip().startswith('La automatización consiste'):
        p.text=p.text.replace('[22]','[21]')

# The measured metrics are already included in each abstract; remove the repeated add-on paragraphs.
for p in list(d.paragraphs):
    if p.text.strip().startswith('Como evidencia de ejecución, el detector YOLOv8n ajustado') or p.text.strip().startswith('As execution evidence, the fine-tuned YOLOv8n detector'):
        p._element.getparent().remove(p._element)
    elif p.text.strip().startswith('Keywords –'):
        p.text=p.text.replace('computer visión','computer vision')

# Merge the two duplicated integration recommendations into one concrete development check.
integration=[p for p in d.paragraphs if p.text.strip().startswith(('Ejecutar una prueba integral reproducible:', 'Completar, en entorno de desarrollo, una prueba integral reproducible:'))]
if integration:
    keep=integration[-1]
    keep.text='Completar en desarrollo una prueba integral reproducible: autenticación, captura/carga, respuesta del detector y OCR, confirmación del operador, endpoint, persistencia SQL y consulta posterior. Registrar versiones, entrada, resultado esperado/obtenido y evidencia; disponer de registro manual ante lecturas ambiguas o indisponibilidad del servicio. Esta verificación técnica no sustituye un piloto de campus.'
    for p in integration[:-1]:
        p._element.getparent().remove(p._element)

# Replace the embedded application screenshot with the privacy-safe crop.
caption=next((p for p in d.paragraphs if p.text.strip().startswith('Figura. Respuesta visual de registro en la interfaz')),None)
if caption:
    previous=caption._p.getprevious()
    while previous is not None and not previous.xpath('.//w:drawing'):
        previous=previous.getprevious()
    if previous is not None:
        previous.getparent().remove(previous)
    safe=Path(r'C:\Users\euco\Desktop\PROYECTO\front\.codex-doc-preview\captura_registro_segura.jpg')
    fig=d.add_paragraph()
    fig.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fig.paragraph_format.space_after=Pt(8)
    fig.add_run().add_picture(str(safe),width=Inches(2.55))
    caption._p.addprevious(fig._p)

# Keep the English abstract and its keywords together instead of splitting the paragraph across pages.
english=next((p for p in d.paragraphs if p.text.strip().startswith('Abstract – The project develops a web prototype')),None)
if english:
    english.paragraph_format.keep_together=True
    english.paragraph_format.keep_with_next=True

# Remove any duplicate copy of the Orozco/Zúñiga bibliographic entry.
orozco=[p for p in d.paragraphs if p.text.strip().startswith('[22] A. F. Orozco Calderón')]
for p in orozco[1:]:
    p._element.getparent().remove(p._element)

# Append the verified institutional thesis used in the revised background.
if not any(p.text.strip().startswith('[22] A. F. Orozco Calderón') for p in d.paragraphs):
    bib=next(p for p in d.paragraphs if p.text.strip().startswith('[21] M. P. Groover'))
    q=d.add_paragraph('[22] A. F. Orozco Calderón and A. A. Zúñiga Rendon, “Desarrollo de un sistema de seguridad y monitoreo que permita la identificación de placas vehiculares mediante visión artificial en tiempo real,” trabajo de titulación, Universidad Politécnica Salesiana, Guayaquil, Ecuador, 2024.')
    bib._p.addnext(q._p)

# Keep one status matrix and remove any duplicate copy from earlier drafts.
matrices=[t for t in d.tables if t.rows and t.cell(0,0).text.strip()=='Componente' and t.cell(0,1).text.strip()=='Estado documentado']
if len(matrices)>1:
    for t in matrices[:-1]:
        t._element.getparent().remove(t._element)
for t in matrices[-1:]:
    for row in t.rows:
        if row.cells[0].text.strip()=='Piloto, tiempos y errores':
            row.cells[1].text='Fuera del alcance actual; requeriría un estudio futuro autorizado.'
        elif row.cells[0].text.strip()=='Prueba de campo y tiempos':
            row.cells[1].text='Fuera del alcance actual; requeriría un estudio futuro autorizado.'
        elif row.cells[0].text.strip()=='Servicio de visión/API':
            row.cells[1].text='Una llamada local de prueba a health/detectar respondió 200; smoke test separado: 7,27 s en CPU para una imagen. No es carga ni flujo completo con frontend y SQL.'

# Preserve the detector evaluation settings needed to interpret and reproduce its metrics.
metric_table=next((t for t in d.tables if t.rows and t.cell(0,0).text.strip()=='Indicador' and t.cell(0,1).text.strip()=='Resultado'),None)
if metric_table and not any(r.cells[0].text.strip()=='Parámetros de evaluación' for r in metric_table.rows):
    for label,value in [('Parámetros de evaluación','640 px; confianza 0,001; IoU 0,70'),('Entorno de evaluación','YOLOv8n ajustado; CUDA 12.8; NVIDIA RTX 3060')]:
        cells=metric_table.add_row().cells
        cells[0].text=label
        cells[1].text=value
        trpr=metric_table.rows[-1]._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))

d.core_properties.title='Protocolo de proyecto de grado con alcance y evidencia revisados'
d.core_properties.subject='Prototipo técnico de detección de placas: resultados y límites documentados'
d.save(f)
print('saved',f,'paragraphs',len(d.paragraphs),'tables',len(d.tables),'bytes',f.stat().st_size)
